"""Export the current 66-slide web presentation to a faithful offline PPTX.

The HTML deck is the editable source. This export captures its designed slides in
presentation order; each PowerPoint slide contains one high-resolution image.
Run from the repository root after starting a local static server there:

    python3 -m http.server 8787 --bind 127.0.0.1
    python export_powerpoint.py --base-url http://127.0.0.1:8787
"""

from __future__ import annotations

import argparse
import json
import tempfile
from pathlib import Path

from PIL import Image, ImageStat
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Inches


HERE = Path(__file__).resolve().parent
PUBLIC_URL = (
    "https://mthoriqzi383.github.io/melooloo-project-hub/"
    "Melooloo_Final_Presentation_Web_20Oct2026/index.html"
)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base-url", default="http://127.0.0.1:8787")
    parser.add_argument(
        "--output", default=str(HERE / "exports" / "Melooloo_Final_Presentation_66.pptx")
    )
    args = parser.parse_args()
    output = Path(args.output).resolve()
    output.parent.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Melooloo — Presentasi Sidang Akhir"
    prs.core_properties.subject = "Versi PowerPoint dari presentasi web 66 slide"
    prs.core_properties.author = "Melooloo"
    blank = prs.slide_layouts[6]
    manifest = []

    with tempfile.TemporaryDirectory(prefix="melooloo-pptx-") as temp_dir:
        with sync_playwright() as pw:
            browser = pw.chromium.launch(channel="chrome", headless=True)
            page = browser.new_page(
                viewport={"width": 1920, "height": 1200},
                device_scale_factor=2,
                reduced_motion="reduce",
            )
            errors = []
            page.on("pageerror", lambda error: errors.append(str(error)))
            page.goto(
                f"{args.base_url.rstrip('/')}/Melooloo_Final_Presentation_Web_20Oct2026/index.html#slide=1",
                wait_until="networkidle",
                timeout=60000,
            )
            route = page.locator(".drawer-item").evaluate_all(
                "items => items.map(item => Number(item.dataset.page))"
            )
            if len(route) != 66 or set(route) != set(range(1, 67)):
                raise RuntimeError(f"Expected all 66 source slides once, got {route}")

            for position, source_number in enumerate(route, start=1):
                if position > 1:
                    page.evaluate("n => { location.hash = `#slide=${n}`; }", source_number)
                page.wait_for_function(
                    "n => document.querySelector('.story-canvas')?.getAttribute('aria-label')?.startsWith(`Slide sumber ${n}:`)",
                    arg=source_number,
                    timeout=30000,
                )
                page.wait_for_function(
                    "() => [...document.querySelectorAll('#storySlide img')].every(img => img.complete && img.naturalWidth > 0)",
                    timeout=30000,
                )
                page.wait_for_load_state("networkidle", timeout=30000)
                # Ensure the slide, rather than a transition overlay, is captured.
                page.locator(".chapter-veil").evaluate("el => { el.hidden = true; }")
                label = page.locator(".story-canvas").get_attribute("aria-label") or ""
                image_path = Path(temp_dir) / f"slide-{position:02d}.jpg"
                page.locator(".stage").screenshot(
                    path=str(image_path), type="jpeg", quality=92, animations="disabled"
                )
                with Image.open(image_path) as image:
                    if image.width < 3000 or image.height < 1700:
                        raise RuntimeError(f"Slide {source_number} exported at insufficient resolution")
                    if max(ImageStat.Stat(image.convert("RGB")).stddev) < 8:
                        raise RuntimeError(f"Slide {source_number} appears blank")

                slide = prs.slides.add_slide(blank)
                slide.shapes.add_picture(
                    str(image_path), 0, 0, width=prs.slide_width, height=prs.slide_height
                )
                title = label.split(": ", 1)[-1]
                slide.notes_slide.notes_text_frame.text = (
                    f"Melooloo presentation {position}/66 · source slide {source_number}\n"
                    f"{title}\n"
                    f"Interactive web slide: {PUBLIC_URL}#slide={source_number}\n"
                    "The web version contains motion, interactive details, and links to the 3D greenhouse."
                )
                manifest.append(
                    {"presentation_slide": position, "source_slide": source_number, "title": title}
                )
                print(f"{position:02d}/66 · source {source_number:02d} · {title}", flush=True)

            browser.close()
            if errors:
                raise RuntimeError("Browser errors during export: " + " | ".join(errors[:5]))

        prs.save(output)

    manifest_path = output.with_suffix(".json")
    manifest_path.write_text(
        json.dumps(
            {
                "format": "faithful image-based PowerPoint export",
                "source": PUBLIC_URL,
                "note": "Edit the HTML deck for element-level changes, then export again.",
                "slides": manifest,
            },
            indent=2,
            ensure_ascii=False,
        )
        + "\n",
        encoding="utf-8",
    )
    print(f"Saved {output} ({output.stat().st_size / 1048576:.1f} MiB)")
    print(f"Saved {manifest_path}")


if __name__ == "__main__":
    main()
